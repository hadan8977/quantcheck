import _test_env  # noqa: F401  -- must stay first: isolates QUANTCHECK_HOME from the real install
import sys
import tempfile
import types
import unittest
from datetime import datetime
from pathlib import Path

sys.modules.setdefault("playwright", types.ModuleType("playwright"))
sys.modules.setdefault(
    "playwright.sync_api",
    types.SimpleNamespace(sync_playwright=lambda: None, TimeoutError=TimeoutError),
)

from openpyxl import load_workbook  # noqa: E402

from quantcheck import picks_excel  # noqa: E402
from quantcheck.diff import compare  # noqa: E402
from quantcheck.historical_resend import _validate_excel_symbols  # noqa: E402

OLD = {
    "fetched_at": "2026-09-30T21:51:44",
    "monthly": {"pick_date": "Updated September 1", "rows": [
        {"symbol": "HPE", "company": "Hewlett Packard Enterprise", "return": "+18.90%"},
        {"symbol": "DELL", "company": "Dell Technologies Inc.", "gt_score": "4.34/5"},
    ]},
    "weekly": {"pick_date": "Updated on Sep 25, 2026", "kind": "watchlist", "rows": [
        {"symbol": "TXG", "company": "10x Genomics, Inc.", "gt_score": "4.79/5"},
    ]},
}
NEW = {
    "fetched_at": "2026-10-01T12:30:00",
    "monthly": {"pick_date": "Updated October 1", "rows": [
        {"symbol": "MRNA", "company": "Moderna, Inc.", "held_since": "2026-10-01", "current_price": "$203.21",
         "buy_or_entry_price": "$186.61", "return": "+8.90%", "gt_score": "4.61/5", "analyst_signal": "Buy +0.42",
         "next_earnings": "Oct 29, 2026", "market_cap": "$81.10B", "revenue_growth_yoy": "-27.62%", "pe_ttm": "—"},
        {"symbol": "DELL", "company": "Dell Technologies Inc.", "gt_score": "3.99/5", "return": "weird value"},
    ]},
    "weekly": {"pick_date": "Updated on Sep 25, 2026", "kind": "watchlist", "rows": [
        {"symbol": "TXG", "company": "10x Genomics, Inc.", "gt_score": "4.79/5", "current_price": "$93.50"},
    ]},
}


def header_index(ws, header_row=picks_excel.HEADER_ROW):
    return {cell.value: cell.column for cell in ws[header_row]}


class PicksExcelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Writing a workbook is the slow part; do it once per variant.
        cls._tmp = tempfile.TemporaryDirectory()
        cls.with_diff = Path(cls._tmp.name) / "with_diff.xlsx"
        cls.without_diff = Path(cls._tmp.name) / "without_diff.xlsx"
        picks_excel.write_report(cls.with_diff, NEW, compare(OLD, NEW), OLD)
        picks_excel.write_report(cls.without_diff, NEW)

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def test_sheets_and_resend_symbol_contract(self):
        # historical_resend relies on the Portfolio / Weekly Watchlist sheets and their Symbol column.
        self.assertEqual(load_workbook(self.with_diff).sheetnames, ["Overview", "Changes", "Portfolio", "Weekly Watchlist"])
        _validate_excel_symbols(self.with_diff, NEW)  # raises on mismatch

        names = load_workbook(self.without_diff).sheetnames
        self.assertNotIn("Changes", names)  # no Changes sheet without a diff
        self.assertEqual(names, ["Overview", "Portfolio", "Weekly Watchlist"])

    def test_values_are_typed_with_number_formats(self):
        ws = load_workbook(self.with_diff)["Portfolio"]
        cols = header_index(ws)
        r = picks_excel.FIRST_DATA_ROW
        self.assertEqual(ws.cell(r, cols["Price"]).value, 203.21)
        self.assertEqual(ws.cell(r, cols["Price"]).number_format, picks_excel.FMT_MONEY)
        self.assertAlmostEqual(ws.cell(r, cols["Return"]).value, 0.089)
        self.assertEqual(ws.cell(r, cols["GT Score"]).value, 4.61)
        self.assertEqual(ws.cell(r, cols["Analyst Consensus"]).value, "Buy")
        self.assertEqual(ws.cell(r, cols["Consensus Score"]).value, 0.42)
        self.assertEqual(ws.cell(r, cols["Next Earnings"]).value, datetime(2026, 10, 29))
        self.assertEqual(ws.cell(r, cols["Market Cap"]).value, 81.10e9)
        self.assertNotIn("P/E (TTM)", cols)  # "—" for every row -> column dropped
        self.assertEqual(ws.cell(r, cols["Status"]).value, "NEW")
        # unparsable values stay visible as text instead of disappearing
        self.assertEqual(ws.cell(r + 1, cols["Return"]).value, "weird value")
        self.assertEqual(ws.cell(r + 1, cols["Status"]).value, "UPDATED")  # gt_score changed

    def test_all_empty_columns_are_dropped(self):
        headers = set(header_index(load_workbook(self.without_diff)["Weekly Watchlist"]))
        self.assertNotIn("Buy Price", headers)
        self.assertIn("Price", headers)

    def test_changes_sheet_lists_added_removed_and_field_updates(self):
        ws = load_workbook(self.with_diff)["Changes"]
        rows = [tuple(c.value for c in r) for r in ws.iter_rows(min_row=picks_excel.FIRST_DATA_ROW)]
        self.assertIn(("Portfolio", "Added", "MRNA", "Moderna, Inc.", None, None, None), rows)
        self.assertIn(("Portfolio", "Removed", "HPE", "Hewlett Packard Enterprise", None, None, None), rows)
        self.assertIn(("Portfolio", "Updated", "DELL", "Dell Technologies Inc.", "GT Score", "4.34/5", "3.99/5"), rows)

    def test_watchlist_sheet_has_signal_and_reason_columns(self):
        data = {
            "fetched_at": "2026-10-07T13:00:00",
            "monthly": {"pick_date": "Updated October 1", "rows": []},
            "weekly": {"pick_date": "Updated on Oct 2, 2026", "kind": "watchlist", "rows": [
                {"symbol": "TEAM", "company": "Atlassian", "current_price": "$196.50", "signal_price": "$187.63",
                 "signal_date": "2026-10-05", "signal_at": "2026-10-05T13:30:00Z", "watch_reason": "Gaining on its sector"},
            ]},
        }
        with tempfile.TemporaryDirectory() as tmp:
            path = picks_excel.write_report(Path(tmp) / "r.xlsx", data)
            ws = load_workbook(path)["Weekly Watchlist"]
            cols = header_index(ws)
            r = picks_excel.FIRST_DATA_ROW
            self.assertEqual(ws.cell(r, cols["Signal Price"]).value, 187.63)
            self.assertEqual(ws.cell(r, cols["Signal Date"]).value, datetime(2026, 10, 5))
            self.assertAlmostEqual(ws.cell(r, cols["Since Signal"]).value, 196.50 / 187.63 - 1)
            self.assertEqual(ws.cell(r, cols["Why Selected"]).value, "Gaining on its sector")


if __name__ == "__main__":
    unittest.main()
