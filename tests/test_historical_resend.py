import _test_env  # noqa: F401  -- must stay first: isolates QUANTCHECK_HOME from the real install
import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from quantcheck.historical_resend import ResendValidationError, execute_resend, main, prepare_resend


def picks(fetched_at, weekly_date, symbols):
    return {
        "fetched_at": fetched_at,
        "source": "https://quantgt.io",
        "monthly": {"pick_date": "Updated on August 1, 2026", "rows": [{"symbol": "DELL"}]},
        "weekly": {
            "pick_date": weekly_date,
            "kind": "watchlist",
            "rows": [{"symbol": symbol} for symbol in symbols],
        },
    }


class HistoricalResendTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        for name in ("state/raw", "output", "screenshots"):
            (self.root / name).mkdir(parents=True)

        old = picks("2026-08-08T16:00:44", "Updated on Jul 31, 2026", ["DELL", "ARM"])
        new = picks("2026-08-09T16:01:01", "Updated on Aug 7, 2026", ["DELL", "FROG"])
        (self.root / "state/raw/picks_raw_2026-08-08_120044.json").write_text(json.dumps(old))
        (self.root / "state/raw/picks_raw_2026-08-09_120101.json").write_text(json.dumps(new))
        workbook = Workbook()
        portfolio = workbook.active
        portfolio.title = "Portfolio"
        for _ in range(5):
            portfolio.append([])
        portfolio.append(["Rank", "Symbol"])
        portfolio.append([1, "DELL"])
        watchlist = workbook.create_sheet("Weekly Watchlist")
        for _ in range(5):
            watchlist.append([])
        watchlist.append(["Rank", "Symbol"])
        watchlist.append([1, "DELL"])
        watchlist.append([2, "FROG"])
        workbook.save(self.root / "output/quantgt_picks_report_2026-08-09_160101.xlsx")
        (self.root / "screenshots/portfolio_2026-08-09_120101.png").write_bytes(b"portfolio")
        (self.root / "screenshots/watchlist_2026-08-09_120101.png").write_bytes(b"watchlist")

    def tearDown(self):
        self.tmp.cleanup()

    def test_selects_snapshot_by_internal_pick_date_reconstructs_diff_and_binds_same_run_attachments(self):
        plan = prepare_resend(self.root, "Updated on Aug 7, 2026")

        self.assertEqual(plan.raw_path.name, "picks_raw_2026-08-09_120101.json")
        self.assertEqual(plan.diff["weekly"]["added"], ["FROG"])
        self.assertEqual(plan.diff["weekly"]["removed"], ["ARM"])
        self.assertIn("+ Added FROG", plan.body)
        self.assertIn("- Removed ARM", plan.body)
        self.assertIn("Delayed alert", plan.body)
        self.assertEqual(plan.subject, "Quant GT · Weekly Watchlist: +FROG -ARM")
        self.assertEqual(plan.summary()["subject"], plan.subject)
        self.assertEqual(
            [path.name for path in plan.attachments],
            [
                "quantgt_picks_report_2026-08-09_160101.xlsx",
                "portfolio_2026-08-09_120101.png",
                "watchlist_2026-08-09_120101.png",
            ],
        )

    def test_fails_closed_when_matching_attachment_is_missing_or_excel_symbols_differ(self):
        shot = self.root / "screenshots/watchlist_2026-08-09_120101.png"
        shot.rename(shot.with_suffix(".hidden"))
        with self.assertRaisesRegex(ResendValidationError, "missing same-run attachment"):
            prepare_resend(self.root, "Updated on Aug 7, 2026")
        shot.with_suffix(".hidden").rename(shot)

        path = self.root / "output/quantgt_picks_report_2026-08-09_160101.xlsx"
        workbook = load_workbook(path)
        for row in workbook["Weekly Watchlist"].iter_rows():
            for cell in row:
                if cell.value == "FROG":
                    cell.value = "WRONG"
        workbook.save(path)
        with self.assertRaisesRegex(ResendValidationError, "Excel symbols do not match raw snapshot"):
            prepare_resend(self.root, "Updated on Aug 7, 2026")

    def test_send_requires_exact_date_confirmation_then_uses_validated_body_attachments_and_private_recipients(self):
        plan = prepare_resend(self.root, "Updated on Aug 7, 2026")
        with patch("quantcheck.historical_resend.send_email_per_recipient") as sender:
            with self.assertRaisesRegex(ResendValidationError, "confirmation date"):
                execute_resend(plan, ["user@example.com"], confirm_date="Updated on Jul 31, 2026")
            sender.assert_not_called()

        with patch("quantcheck.historical_resend.send_email_per_recipient", return_value=(["user@example.com"], [])) as sender:
            delivered, failed = execute_resend(plan, ["user@example.com"], confirm_date="Updated on Aug 7, 2026")

        self.assertEqual((delivered, failed), (["user@example.com"], []))
        sender.assert_called_once_with(
            plan.subject,
            plan.body,
            to=["user@example.com"],
            attachments=list(plan.attachments),
            html=plan.html,
            retries=2,
        )

    def test_cli_recipient_flag_is_admin_preview_otherwise_the_full_subscriber_route(self):
        # --recipient: --send must reach exactly the given address and never call subscriber_recipients
        # (the real subscriber list), even though .env is loaded. Without it: backward-compatible subscriber route.
        (self.root / ".env").write_text("NOTIFY_EMAIL_TO=real-subscriber@example.com\n", encoding="utf-8")
        base = ["historical_resend", "--weekly-date", "Updated on Aug 7, 2026", "--root", str(self.root),
                "--send", "--confirm-date", "Updated on Aug 7, 2026"]
        cases = {
            "explicit recipient": (["--recipient", "admin@example.com"], ["admin@example.com"], "explicit_recipient"),
            "subscriber route": ([], ["real-subscriber@example.com"], "subscriber_route"),
        }
        for name, (extra, expected_to, source) in cases.items():
            with self.subTest(name):
                buf = io.StringIO()
                with patch("quantcheck.historical_resend.send_email_per_recipient", return_value=(expected_to, [])) as sender, \
                     patch("sys.argv", base + extra), patch.dict("os.environ", {}, clear=True), redirect_stdout(buf):
                    if extra:
                        with patch("quantcheck.historical_resend.subscriber_recipients") as subscriber_route:
                            main()
                        subscriber_route.assert_not_called()
                    else:
                        main()

                sender.assert_called_once()
                self.assertEqual(sender.call_args.kwargs["to"], expected_to)
                summary = json.loads(buf.getvalue())
                self.assertEqual(summary["recipients_source"], source)
                self.assertEqual(summary["delivered"], len(expected_to))

    def test_unknown_weekly_date_fails_closed(self):
        with self.assertRaisesRegex(ResendValidationError, "no raw snapshot has weekly.pick_date"):
            prepare_resend(self.root, "Updated on Jan 1, 2026")

    def test_duplicate_snapshots_of_the_same_update_still_reconstruct_the_diff(self):
        duplicate = picks("2026-08-10T13:00:43", "Updated on Aug 7, 2026", ["DELL", "FROG"])
        (self.root / "state/raw/picks_raw_2026-08-10_090043.json").write_text(json.dumps(duplicate))

        plan = prepare_resend(self.root, "Updated on Aug 7, 2026")
        self.assertTrue(plan.diff["changed"])
        self.assertIn("Changes:", plan.body)


if __name__ == "__main__":
    unittest.main()
