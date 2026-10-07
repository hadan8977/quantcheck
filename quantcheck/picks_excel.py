"""Excel report attached to every picks alert.

Sheets: Overview, Changes (only when there is a diff), Portfolio, and the
weekly list ("Weekly Watchlist" / "Weekly Picks").

Values are written as real numbers/dates with number formats (prices,
percentages, GT Score, signal score, earnings dates), so subscribers can sort
and filter; anything that does not parse is kept as the original text.

Contract relied on by quantcheck.historical_resend._validate_excel_symbols:
the Portfolio and weekly sheets keep those exact sheet names, have a header
cell exactly "Symbol", and below that header the Symbol column contains
nothing but ticker symbols.

Styling follows the existing house style: white background, no row banding,
green header accents.
"""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path
from typing import Any, Callable, Dict, List

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from quantcheck import picks_email
from quantcheck.picks_format import (
    display_date,
    field_label,
    fetched_at_utc,
    format_fetched,
    is_blank,
    parse_date,
    parse_gt_score,
    parse_money,
    parse_pct,
    parse_signal,
    rows_by_symbol,
    section_title,
    since_signal,
    stock_count,
    NY,
)

GREEN = "16A34A"
GREEN_DARK = "0F7A36"
GREEN_SOFT = "DCFCE7"
RED = "B91C1C"
RED_SOFT = "FEE2E2"
UP = "15803D"
SIGNED_FORMATS = {"+0.00%;-0.00%;0.00%", "+0.00;-0.00;0.00"}
AMBER = "B45309"
TEXT = "0F172A"
MUTED = "64748B"
GRID = "E2E8F0"
FONT = "Aptos"

HEADER_ROW = 4
FIRST_DATA_ROW = HEADER_ROW + 1
EARNINGS_SOON_DAYS = 14

FMT_MONEY = '$#,##0.00'
FMT_PCT = '+0.00%;-0.00%;0.00%'
FMT_GT = '0.00'
FMT_SCORE = '+0.00;-0.00;0.00'
FMT_DATE = 'mmm d, yyyy'
FMT_CAP = '[>=1000000000]$#,##0.00,,,"B";[>=1000000]$#,##0.00,,"M";$#,##0'


def _signal_label(value: Any) -> Any:
    label, _ = parse_signal(value)
    return label or None


def _signal_score(value: Any) -> Any:
    return parse_signal(value)[1]


def _number(value: Any) -> Any:
    return value if isinstance(value, (int, float)) else None


# (header, field, parser, number_format, width, align)
Column = tuple[str, str, Callable[[Any], Any], str | None, float, str]

PORTFOLIO_COLUMNS: List[Column] = [
    ("Symbol", "symbol", str, None, 10, "left"),
    ("Company", "company", str, None, 30, "left"),
    ("Sector", "sector", str, None, 22, "left"),
    ("Entry Date", "held_since", parse_date, FMT_DATE, 13, "center"),
    ("Entry Price", "buy_or_entry_price", parse_money, FMT_MONEY, 12, "right"),
    ("Price", "current_price", parse_money, FMT_MONEY, 12, "right"),
    ("Return", "return", parse_pct, FMT_PCT, 11, "right"),
    ("GT Score", "gt_score", parse_gt_score, FMT_GT, 10, "center"),
    ("Analyst Consensus", "analyst_signal", _signal_label, None, 16, "left"),
    ("Consensus Score", "analyst_signal", _signal_score, FMT_SCORE, 14, "right"),
    ("Next Earnings", "next_earnings", parse_date, FMT_DATE, 14, "center"),
    ("Market Cap", "market_cap", parse_money, FMT_CAP, 13, "right"),
    ("Revenue (TTM)", "revenue_ttm", parse_money, FMT_CAP, 14, "right"),
    ("Revenue Growth", "revenue_growth_yoy", parse_pct, FMT_PCT, 15, "right"),
    ("P/E (TTM)", "pe_ttm", parse_money, '0.0', 10, "right"),
]

WEEKLY_COLUMNS: List[Column] = [
    ("Symbol", "symbol", str, None, 10, "left"),
    ("Company", "company", str, None, 30, "left"),
    ("Sector", "sector", str, None, 22, "left"),
    ("GT Score", "gt_score", parse_gt_score, FMT_GT, 10, "center"),
    ("Price", "current_price", parse_money, FMT_MONEY, 12, "right"),
    ("Buy Price", "buy_or_entry_price", parse_money, FMT_MONEY, 12, "right"),
    ("Signal Date", "signal_date", parse_date, FMT_DATE, 13, "center"),
    ("Signal Price", "signal_price", parse_money, FMT_MONEY, 12, "right"),
    ("Since Signal", "_since_signal", _number, FMT_PCT, 12, "right"),
    ("Why Selected", "watch_reason", str, None, 24, "left"),
    ("Analyst Consensus", "analyst_signal", _signal_label, None, 16, "left"),
    ("Consensus Score", "analyst_signal", _signal_score, FMT_SCORE, 14, "right"),
    ("Next Earnings", "next_earnings", parse_date, FMT_DATE, 14, "center"),
    ("Market Cap", "market_cap", parse_money, FMT_CAP, 13, "right"),
    ("Revenue (TTM)", "revenue_ttm", parse_money, FMT_CAP, 14, "right"),
    ("Revenue Growth", "revenue_growth_yoy", parse_pct, FMT_PCT, 15, "right"),
    ("P/E (TTM)", "pe_ttm", parse_money, '0.0', 10, "right"),
]

THIN = Side(style="thin", color=GRID)


def _font(**kwargs: Any) -> Font:
    return Font(name=FONT, size=kwargs.pop("size", 10), color=kwargs.pop("color", TEXT), **kwargs)


def _cell_value(raw: Any, parser: Callable[[Any], Any]) -> Any:
    if is_blank(raw):
        return None
    if parser is str:
        return str(raw)
    parsed = parser(raw)
    return parsed if parsed is not None else str(raw)  # keep unparsable text visible


def _title(ws, title: str, subtitle: str, last_col: int) -> None:
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = _font(size=18, bold=True)
    ws.row_dimensions[1].height = 28
    ws["A2"] = subtitle
    ws["A2"].font = _font(size=10, color=MUTED)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(last_col, 4))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(last_col, 4))


def _header(ws, row: int, headers: List[str], start_col: int = 1) -> None:
    for offset, text in enumerate(headers):
        cell = ws.cell(row, start_col + offset, text)
        cell.font = _font(bold=True, color=GREEN_DARK)
        cell.fill = PatternFill("solid", fgColor=GREEN_SOFT)
        cell.border = Border(bottom=Side(style="medium", color=GREEN))
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 22


def _body_cell(cell, align: str = "left", number_format: str | None = None, **font: Any) -> None:
    # Colour signed numbers explicitly: [Color10]-style formats render differently in
    # Excel, LibreOffice, WPS and mobile viewers.
    if number_format in SIGNED_FORMATS and isinstance(cell.value, (int, float)) and "color" not in font:
        font["color"] = UP if cell.value > 0 else (RED if cell.value < 0 else TEXT)
        font.setdefault("bold", True)
    cell.font = _font(**font)
    cell.alignment = Alignment(horizontal=align, vertical="center", indent=0 if align == "center" else 1)
    cell.border = Border(bottom=THIN)
    if number_format:
        cell.number_format = number_format


def _print_setup(ws, last_col: int, last_row: int, header_row: int) -> None:
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.print_area = f"A1:{get_column_letter(last_col)}{max(last_row, header_row)}"


def _report_date(data: Dict[str, Any]) -> date:
    fetched = fetched_at_utc(data)
    return fetched.astimezone(NY).date() if fetched else date.today()


def _write_picks_sheet(wb, key: str, data: Dict[str, Any], diff: Dict[str, Any] | None) -> None:
    section = data.get(key) or {}
    as_of = fetched_at_utc(data)
    rows = [{**row, "_since_signal": since_signal(row, as_of)} for row in section.get("rows") or []]
    title = section_title(key, section)
    columns = PORTFOLIO_COLUMNS if key == "monthly" else WEEKLY_COLUMNS
    # Drop columns Quant GT left empty for every row (e.g. watchlist Buy Price).
    columns = [c for c in columns if c[1] == "symbol" or any(not is_blank(r.get(c[1])) for r in rows)]
    added = set(((diff or {}).get(key) or {}).get("added") or [])
    changed_fields = {
        item.get("symbol"): set((item.get("fields") or {}).keys())
        for item in ((diff or {}).get(key) or {}).get("changed") or []
    }

    ws = wb.create_sheet(title)
    headers = ["#", *[c[0] for c in columns], "Status"]
    last_col = len(headers)
    _title(ws, title, f"Updated {display_date(section.get('pick_date'))} · {stock_count(len(rows))}", last_col)
    _header(ws, HEADER_ROW, headers)
    ws.column_dimensions["A"].width = 5
    for idx, column in enumerate(columns, 2):
        ws.column_dimensions[get_column_letter(idx)].width = column[4]
    ws.column_dimensions[get_column_letter(last_col)].width = 12

    soon = _report_date(data) + timedelta(days=EARNINGS_SOON_DAYS)
    for offset, row in enumerate(rows):
        r = FIRST_DATA_ROW + offset
        _body_cell(ws.cell(r, 1, offset + 1), "center", color=MUTED)
        for idx, (header, field, parser, fmt, _width, align) in enumerate(columns, 2):
            value = _cell_value(row.get(field), parser)
            cell = ws.cell(r, idx, value)
            numeric = not isinstance(value, str)
            _body_cell(cell, align, fmt if numeric else None)
            if field == "symbol":
                cell.font = _font(bold=True, color=GREEN_DARK)
            elif field == "next_earnings" and isinstance(value, date) and value <= soon:
                cell.font = _font(bold=True, color=AMBER)
            elif field in changed_fields.get(row.get("symbol"), set()):
                keep = cell.font.color.rgb if cell.font.color is not None and isinstance(cell.font.color.rgb, str) else "FF" + TEXT
                cell.font = _font(bold=True, color=keep[-6:])
        status = "NEW" if row.get("symbol") in added else ("UPDATED" if changed_fields.get(row.get("symbol")) else None)
        status_cell = ws.cell(r, last_col, status)
        _body_cell(status_cell, "center", bold=bool(status), color=GREEN_DARK if status == "NEW" else MUTED)
        if status == "NEW":
            status_cell.fill = PatternFill("solid", fgColor=GREEN_SOFT)

    last_row = FIRST_DATA_ROW + len(rows) - 1
    if rows:
        gt_col = next((i for i, c in enumerate(columns, 2) if c[1] == "gt_score"), None)
        if gt_col:
            letter = get_column_letter(gt_col)
            ws.conditional_formatting.add(
                f"{letter}{FIRST_DATA_ROW}:{letter}{last_row}",
                ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="BBF7D0"),
            )
        ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(last_col)}{last_row}"
    ws.freeze_panes = ws.cell(FIRST_DATA_ROW, 3)  # keep # and Symbol visible when scrolling
    note_row = max(last_row, HEADER_ROW) + 2
    note = ws.cell(note_row, 3, f"Next Earnings in amber = within {EARNINGS_SOON_DAYS} days. Status: NEW = just added, UPDATED = a tracked field changed (bold).")
    note.font = _font(size=9, color=MUTED, italic=True)
    _print_setup(ws, last_col, last_row, HEADER_ROW)


def _write_overview(ws, data: Dict[str, Any], diff: Dict[str, Any] | None) -> None:
    ws.title = "Overview"
    fetched = format_fetched(data)
    _title(ws, "Quant GT Picks Report", f"Data checked {fetched}" if fetched else "", 9)
    if diff is not None:
        headline = picks_email.build_subject(diff, data).removeprefix(picks_email.SUBJECT_PREFIX)
        ws["A3"] = f"This update: {headline}"
        ws["A3"].font = _font(size=11, bold=True, color=GREEN_DARK)
        ws.merge_cells("A3:I3")

    _header(ws, 5, ["List", "Updated", "Stocks", "Changes"])
    for offset, key in enumerate(("monthly", "weekly")):
        r = 6 + offset
        section = data.get(key) or {}
        sec_diff = ((diff or {}).get(key) or {}) if diff else {}
        moves = " ".join([*(f"+{s}" for s in sec_diff.get("added") or []), *(f"-{s}" for s in sec_diff.get("removed") or [])])
        n_changed = len(sec_diff.get("changed") or [])
        change_text = moves or (f"{n_changed} stock(s) updated" if n_changed else ("—" if diff is None else "No changes"))
        values = [section_title(key, section), display_date(section.get("pick_date")), len(section.get("rows") or []), change_text]
        for c, value in enumerate(values, 1):
            _body_cell(ws.cell(r, c, value), "center" if c == 3 else "left", bold=(c == 1))
    ws.merge_cells("D5:I5")
    for r in (6, 7):
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=9)

    # Side-by-side quick lists.
    top = 10
    blocks = [
        ("monthly", 1, [("Symbol", "symbol", str, None), ("Company", "company", str, None), ("Return", "return", parse_pct, FMT_PCT), ("GT Score", "gt_score", parse_gt_score, FMT_GT)]),
        ("weekly", 6, [("Symbol", "symbol", str, None), ("Company", "company", str, None), ("GT Score", "gt_score", parse_gt_score, FMT_GT), ("Consensus", "analyst_signal", str, None)]),
    ]
    for key, start_col, cols in blocks:
        section = data.get(key) or {}
        label = ws.cell(top - 1, start_col, f"{section_title(key, section)} · {display_date(section.get('pick_date'))}")
        label.font = _font(size=11, bold=True)
        _header(ws, top, [c[0] for c in cols], start_col)
        added = set(((diff or {}).get(key) or {}).get("added") or [])
        for offset, row in enumerate(section.get("rows") or []):
            r = top + 1 + offset
            for idx, (header, field, parser, fmt) in enumerate(cols):
                value = _cell_value(row.get(field), parser)
                cell = ws.cell(r, start_col + idx, value)
                _body_cell(cell, "left" if idx < 2 else "right", fmt if not isinstance(value, str) else None)
                if field == "symbol":
                    cell.font = _font(bold=True, color=GREEN_DARK)
                    if row.get("symbol") in added:
                        cell.value = f"{row.get('symbol')}  NEW"
    for col, width in {1: 18, 2: 26, 3: 11, 4: 10, 5: 3, 6: 10, 7: 26, 8: 10, 9: 18}.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    last_row = top + max(len((data.get(k) or {}).get("rows") or []) for k in ("monthly", "weekly"))
    _print_setup(ws, 9, last_row, top)


def _write_changes(wb, data: Dict[str, Any], diff: Dict[str, Any], previous: Dict[str, Any] | None) -> None:
    ws = wb.create_sheet("Changes")
    headers = ["List", "Change", "Symbol", "Company", "Field", "Previous", "New"]
    _title(ws, "What changed", picks_email.build_subject(diff, data).removeprefix(picks_email.SUBJECT_PREFIX), len(headers))
    _header(ws, HEADER_ROW, headers)
    rows: List[tuple[Any, ...]] = []
    for key in ("monthly", "weekly"):
        sec = (diff.get(key) or {}) if diff.get(key, {}).get("changed_flag") else {}
        if not sec:
            continue
        title = section_title(key, data.get(key))
        current = rows_by_symbol((data.get(key) or {}).get("rows"))
        prior = rows_by_symbol(((previous or {}).get(key) or {}).get("rows"))
        if sec.get("date"):
            rows.append((title, "Update date", "", "", "", display_date(sec["date"].get("old")), display_date(sec["date"].get("new"))))
        for symbol in sec.get("added") or []:
            rows.append((title, "Added", symbol, current.get(symbol, {}).get("company", ""), "", "", ""))
        for symbol in sec.get("removed") or []:
            rows.append((title, "Removed", symbol, prior.get(symbol, {}).get("company", ""), "", "", ""))
        for item in sec.get("changed") or []:
            symbol = item.get("symbol") or ""
            for name, vals in (item.get("fields") or {}).items():
                if name in picks_email.HIDDEN_FIELDS:
                    continue
                rows.append((title, "Updated", symbol, current.get(symbol, {}).get("company", ""), field_label(name), vals.get("old"), vals.get("new")))
    tone = {"Added": (GREEN_DARK, GREEN_SOFT), "Removed": (RED, RED_SOFT)}
    for offset, values in enumerate(rows):
        r = FIRST_DATA_ROW + offset
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c, value if not is_blank(value) else None)
            _body_cell(cell, "left", bold=(c == 3))
        if values[1] in tone:
            fg, bg = tone[values[1]]
            change_cell = ws.cell(r, 2)
            change_cell.font = _font(bold=True, color=fg)
            change_cell.fill = PatternFill("solid", fgColor=bg)
    for col, width in enumerate([18, 13, 10, 30, 16, 20, 20], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(FIRST_DATA_ROW, 1)
    last_row = FIRST_DATA_ROW + len(rows) - 1
    if rows:
        ws.auto_filter.ref = f"A{HEADER_ROW}:G{last_row}"
    _print_setup(ws, len(headers), last_row, HEADER_ROW)


def write_report(path: Path, data: Dict[str, Any], diff: Dict[str, Any] | None = None, previous: Dict[str, Any] | None = None) -> Path:
    wb = Workbook()
    _write_overview(wb.active, data, diff)
    if diff is not None and diff.get("changed"):
        _write_changes(wb, data, diff, previous)
    _write_picks_sheet(wb, "monthly", data, diff)
    _write_picks_sheet(wb, "weekly", data, diff)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
