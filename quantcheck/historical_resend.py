from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import datetime
import json
import os
from pathlib import Path
from typing import Any, Sequence

from openpyxl import load_workbook

from quantcheck.config import load_env
from quantcheck.diff import compare
from quantcheck.gmail_api_notify import send_email_per_recipient
from quantcheck.notify_routes import subscriber_recipients
from quantcheck.picks_check import build_notification_body, build_notification_html, strip_dynamic


class ResendValidationError(RuntimeError):
    pass


@dataclass(frozen=True)
class ResendPlan:
    target_weekly_date: str
    raw_path: Path
    previous_raw_path: Path
    attachments: tuple[Path, Path, Path]
    data: dict[str, Any]
    diff: dict[str, Any]
    body: str
    html: str

    def summary(self) -> dict[str, Any]:
        weekly = self.diff.get("weekly", {})
        return {
            "mode": "preview",
            "target_weekly_date": self.target_weekly_date,
            "raw": str(self.raw_path),
            "previous_raw": str(self.previous_raw_path),
            "attachments": [str(path) for path in self.attachments],
            "added": weekly.get("added", []),
            "removed": weekly.get("removed", []),
            "changed_rows": len(weekly.get("changed", [])),
        }


def _load_snapshots(root: Path) -> list[tuple[datetime, Path, dict[str, Any]]]:
    snapshots = []
    for path in (root / "state/raw").glob("picks_raw_*.json"):
        if "_test_" in path.name:
            continue
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            fetched_at = datetime.fromisoformat(str(data["fetched_at"]).replace("Z", "+00:00"))
        except (KeyError, TypeError, ValueError, json.JSONDecodeError):
            continue
        snapshots.append((fetched_at, path, data))
    return sorted(snapshots, key=lambda item: (item[0], item[1].name))


def _same_run_attachments(root: Path, raw_path: Path, data: dict[str, Any]) -> tuple[Path, Path, Path]:
    ny_stamp = raw_path.stem.removeprefix("picks_raw_")
    try:
        fetched = datetime.fromisoformat(str(data["fetched_at"]).replace("Z", "+00:00"))
    except (KeyError, TypeError, ValueError) as exc:
        raise ResendValidationError("selected raw snapshot has invalid fetched_at") from exc
    utc_stamp = fetched.strftime("%Y-%m-%d_%H%M%S")
    attachments = (
        root / "output" / f"quantgt_picks_report_{utc_stamp}.xlsx",
        root / "screenshots" / f"portfolio_{ny_stamp}.png",
        root / "screenshots" / f"watchlist_{ny_stamp}.png",
    )
    missing = [str(path) for path in attachments if not path.is_file() or path.stat().st_size == 0]
    if missing:
        raise ResendValidationError("missing same-run attachment: " + ", ".join(missing))
    _validate_excel_symbols(attachments[0], data)
    return attachments


def _symbols(rows: Sequence[dict[str, Any]]) -> set[str]:
    return {str(row.get("symbol", "")).strip().upper() for row in rows if str(row.get("symbol", "")).strip()}


def _sheet_symbols(path: Path, sheet_name: str) -> set[str]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        for header in rows:
            symbol_columns = [index for index, value in enumerate(header) if str(value).strip() == "Symbol"]
            if symbol_columns:
                symbol_column = symbol_columns[0]
                break
        else:
            raise ResendValidationError(f"cannot validate Excel sheet {sheet_name!r}: Symbol header not found")
        return {str(row[symbol_column]).strip().upper() for row in rows if len(row) > symbol_column and row[symbol_column]}
    except ResendValidationError:
        raise
    except (KeyError, ValueError, OSError) as exc:
        raise ResendValidationError(f"cannot validate Excel sheet {sheet_name!r}: {exc}") from exc


def _validate_excel_symbols(path: Path, data: dict[str, Any]) -> None:
    expected = {
        "Portfolio": _symbols(data.get("monthly", {}).get("rows", [])),
        "Weekly Watchlist": _symbols(data.get("weekly", {}).get("rows", [])),
    }
    actual = {sheet_name: _sheet_symbols(path, sheet_name) for sheet_name in expected}
    mismatches = [
        f"{sheet_name}: expected={sorted(expected[sheet_name])} actual={sorted(actual[sheet_name])}"
        for sheet_name in expected
        if expected[sheet_name] != actual[sheet_name]
    ]
    if mismatches:
        raise ResendValidationError("Excel symbols do not match raw snapshot: " + "; ".join(mismatches))


def prepare_resend(root: Path, target_weekly_date: str) -> ResendPlan:
    root = Path(root)
    snapshots = _load_snapshots(root)
    matches = [index for index, (_, _, data) in enumerate(snapshots) if data.get("weekly", {}).get("pick_date") == target_weekly_date]
    if not matches:
        raise ResendValidationError(f"no raw snapshot has weekly.pick_date={target_weekly_date!r}")

    selected_index = matches[0]
    if selected_index == 0:
        raise ResendValidationError("no previous raw snapshot exists before target update")
    _, raw_path, data = snapshots[selected_index]
    _, previous_raw_path, previous = snapshots[selected_index - 1]

    diff = compare(strip_dynamic(previous), strip_dynamic(data))
    weekly_diff = diff.get("weekly", {})
    if not diff.get("changed") or not weekly_diff.get("changed_flag"):
        raise ResendValidationError("target update has no reconstructed Weekly Watchlist diff")
    if weekly_diff.get("date", {}).get("new") != target_weekly_date:
        raise ResendValidationError("reconstructed diff does not end at target weekly date")

    attachments = _same_run_attachments(root, raw_path, data)
    context = f"historical resend: {target_weekly_date}"
    body = build_notification_body(data, diff, context=context)
    html = build_notification_html(data, diff, context=context)
    required_fragments = ["Changes:", target_weekly_date]
    required_fragments.extend(weekly_diff.get("added", []))
    required_fragments.extend(weekly_diff.get("removed", []))
    missing_fragments = [fragment for fragment in required_fragments if fragment and fragment not in body]
    if missing_fragments:
        raise ResendValidationError("rendered body is missing change details: " + ", ".join(missing_fragments))

    return ResendPlan(
        target_weekly_date=target_weekly_date,
        raw_path=raw_path,
        previous_raw_path=previous_raw_path,
        attachments=attachments,
        data=data,
        diff=diff,
        body=body,
        html=html,
    )


def execute_resend(plan: ResendPlan, recipients: Sequence[str], confirm_date: str) -> tuple[list[str], list[str]]:
    if confirm_date != plan.target_weekly_date:
        raise ResendValidationError("confirmation date must exactly match the selected Weekly Watchlist date")
    normalized = [recipient.strip() for recipient in recipients if recipient.strip()]
    if not normalized or len(normalized) != len({recipient.lower() for recipient in normalized}):
        raise ResendValidationError("recipient list is empty or contains duplicates")
    delivered, failed = send_email_per_recipient(
        "Quant GT Picks Updated",
        plan.body,
        to=normalized,
        attachments=list(plan.attachments),
        html=plan.html,
        retries=2,
    )
    if failed or len(delivered) != len(normalized):
        raise ResendValidationError(f"incomplete resend: delivered={len(delivered)} expected={len(normalized)} failed={failed}")
    return delivered, failed


def main() -> None:
    parser = argparse.ArgumentParser(description="Fail-closed historical Quantcheck report resend")
    parser.add_argument("--weekly-date", required=True, help="Exact internal date, e.g. 'Updated on Aug 7, 2026'")
    parser.add_argument("--root", type=Path, default=Path("/opt/quantcheck"))
    parser.add_argument("--send", action="store_true", help="Send after all validation gates pass")
    parser.add_argument("--confirm-date", default="", help="Must exactly match --weekly-date when --send is used")
    args = parser.parse_args()

    plan = prepare_resend(args.root, args.weekly_date)
    summary = plan.summary()
    if not args.send:
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        return

    load_env(args.root)
    recipients = subscriber_recipients(dict(os.environ))
    delivered, failed = execute_resend(plan, recipients, args.confirm_date)
    summary.update({"mode": "sent", "delivered": len(delivered), "failed": failed})
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
